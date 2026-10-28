import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_document():
    doc = docx.Document()
    
    # 1. Global Style Definitions
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    
    # Paragraph spacing
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15
    
    # 2. Configure Header
    section = doc.sections[0]
    section.top_margin = Inches(1.8)
    header = section.header
    
    h_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    h_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    header_img_path = os.path.join(os.getcwd(), 'header.png')
    if os.path.exists(header_img_path):
        run_img = h_para.add_run()
        run_img.add_picture(header_img_path, width=Inches(6.2))
    else:
        h_para.add_run("Header image not found.")
    
    # --- HELPER FUNCTIONS ---
    def add_justified_para(text, style=None):
        p = doc.add_paragraph(text, style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def add_section_heading(text):
        doc.add_paragraph() # Spacing
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.underline = True
        r.font.size = Pt(13)
        p.paragraph_format.space_after = Pt(4)

    def add_meta_field(label, value):
        p = doc.add_paragraph()
        r1 = p.add_run(f'{label}: ')
        r1.bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(4)

    # --- CONTENT ---
    doc.add_paragraph() # Spacing
    
    # Title
    proposal_title = doc.add_paragraph()
    proposal_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = proposal_title.add_run('SH-26 MAJOR PROJECT PROPOSAL SUBMISSION')
    run_title.bold = True
    run_title.underline = True
    run_title.font.size = Pt(14)
    
    doc.add_paragraph() # Spacing
    
    add_meta_field('Project Title', 'Stock Price Prediction System: AI-Powered Financial Analytics & Recommendation Platform')
    add_meta_field('Project Domain', 'Artificial Intelligence & Machine Learning')
    add_meta_field('Project Type', 'Software')
    add_meta_field('Project Category', 'FinTech / Quantitative Finance')
    
    sdg_p = doc.add_paragraph()
    r_sdg = sdg_p.add_run('SDG Mapping & Justification:\n')
    r_sdg.bold = True
    
    b_sdg = doc.add_paragraph('SDG 8 – Decent Work and Economic Growth', style='List Bullet')
    b_sdg.runs[0].underline = True
    b_sdg.paragraph_format.space_after = Pt(2)
    
    add_justified_para('The Stock Price Prediction System democratizes algorithmic trading and quantitative finance for retail investors. By providing accessible, institutional-grade financial analytics, it promotes financial literacy and smarter investment decisions, thereby supporting economic growth and wealth generation.')
    
    add_meta_field('Project Guide', '[Name of the Project Guide]')
    
    # Team Details
    doc.add_paragraph().add_run('Team Details:').bold = True
    
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'SrNo.'
    hdr_cells[1].text = 'Student ID'
    hdr_cells[2].text = 'Student Name'
    for cell in hdr_cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    for i in range(1, 5):
        row_cells = table.rows[i].cells
        row_cells[0].text = str(i)
        row_cells[1].text = f'[ID {i}]'
        row_cells[2].text = f'[Name {i}]'
        for cell in row_cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_section_heading("Abstract:")
    add_justified_para("The Stock Price Prediction System is an AI-powered, full-stack web application designed for real-time stock price analysis. Unlike traditional stock charting tools that require manual interpretation, this platform acts as an intelligent quantitative assistant. It forecasts future stock trends using a custom deep learning meta-ensemble (LSTM, GRU, and Conv1D) and provides actionable BUY/SELL/HOLD recommendations by combining technical indicators and NLP-based news sentiment analysis into a single unified platform.")
    add_justified_para("Users can search for global or Indian (NSE/BSE) stocks, select timeframes, and instantly receive customized technical analysis, price forecasting, and market sentiment, enabling data-driven investment decisions.")

    add_section_heading("Problem Statement")
    add_justified_para("Retail investors often struggle to analyze massive amounts of market data, technical indicators, and financial news simultaneously. Existing platforms either focus solely on raw charting or offer overly complex, expensive institutional tools. There is a lack of accessible platforms that seamlessly integrate machine learning price forecasts, automated technical analysis, and real-time news sentiment. This creates a need for an intuitive, centralized intelligent platform that democratizes quantitative finance and simplifies the trading process.")

    add_section_heading("Problem Solution")
    add_justified_para("The Stock Price Prediction System is an AI-powered financial platform that generates real-time price predictions based on historical OHLCV data, macroeconomic indicators, and technical parameters. It combines a robust REST API backend with an interactive Streamlit frontend to deliver automated technical analysis, sentiment scoring, and AI-driven recommendations in a single, user-friendly dashboard.")

    add_section_heading("Objectives")
    doc.add_paragraph("To develop a decoupled full-stack AI-powered stock prediction platform.", style='List Bullet')
    doc.add_paragraph("To implement a Meta-Ensemble Deep Learning model for accurate time-series price forecasting.", style='List Bullet')
    doc.add_paragraph("To generate automated BUY/SELL/HOLD recommendations using a multi-factor scoring system.", style='List Bullet')
    doc.add_paragraph("To integrate real-time NLP sentiment analysis from financial news.", style='List Bullet')
    doc.add_paragraph("To calculate and visualize advanced technical indicators (MA, RSI, MACD, Bollinger Bands) instantly.", style='List Bullet')

    doc.add_page_break()

    add_section_heading("Scope of the Project")
    doc.add_paragraph().add_run("Features included").bold = True
    doc.add_paragraph("AI Price Predictor utilizing a fused Deep Learning model (LSTM + Attention + GRU + Conv1D).", style='List Bullet')
    doc.add_paragraph("Smart Recommendation Engine that aggregates Technicals, AI predictions, and Sentiment.", style='List Bullet')
    doc.add_paragraph("Advanced Technical Analysis with interactive Plotly charting.", style='List Bullet')
    doc.add_paragraph("Real-Time Data Fetching for both US and Indian stock markets.", style='List Bullet')
    doc.add_paragraph("Live News Sentiment Analysis using FinBERT.", style='List Bullet')
    doc.add_paragraph("Premium Dashboard featuring a dark-themed UI.", style='List Bullet')

    doc.add_paragraph().add_run("Features Excluded").bold = True
    doc.add_paragraph("Live trade execution (Broker integration)", style='List Bullet')
    doc.add_paragraph("High-Frequency Trading (HFT) capabilities", style='List Bullet')
    doc.add_paragraph("Paper trading portfolio management", style='List Bullet')

    doc.add_paragraph().add_run("Future Scope").bold = True
    doc.add_paragraph("Paper Trading Integration for tracking recommendation performance.", style='List Bullet')
    doc.add_paragraph("Alternative Data integration (e.g., social media sentiment via Twitter/Reddit).", style='List Bullet')
    doc.add_paragraph("Probabilistic forecasting using Bayesian Neural Networks.", style='List Bullet')

    add_section_heading("Proposed System")
    doc.add_paragraph().add_run("Proposed Solution").bold = True
    add_justified_para("The system employs a microservices architecture isolating the UI from the heavy machine learning workloads. It integrates live market data fetching, vectorized indicator calculation, NLP sentiment scoring, and deep learning inference into a single, cohesive workflow accessible via a REST API.")

    doc.add_paragraph().add_run("Architecture").bold = True
    doc.add_paragraph("Frontend (Streamlit): Serves as the presentation layer.", style='List Bullet')
    doc.add_paragraph("Backend (Flask): Acts as the REST API gateway.", style='List Bullet')
    doc.add_paragraph("Services Layer: Contains StockService, IndicatorService, PredictionService, and RecommendationService.", style='List Bullet')
    doc.add_paragraph("AI Engine: Deep Learning ensemble deployed locally, caching models for O(1) subsequent inference times.", style='List Bullet')

    arch_img_path = os.path.join(os.getcwd(), 'architecture.png')
    if os.path.exists(arch_img_path):
        doc.add_paragraph() # Spacing
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(arch_img_path, width=Inches(6.0))

    doc.add_paragraph().add_run("Advantages").bold = True
    doc.add_paragraph("Accessible institutional-grade analytics.", style='List Bullet')
    doc.add_paragraph("Highly accurate predictions through Meta-Ensemble deep learning.", style='List Bullet')
    doc.add_paragraph("Real-time insights and live market fetching.", style='List Bullet')
    doc.add_paragraph("Time-saving automated technical analysis.", style='List Bullet')

    doc.add_page_break()

    add_section_heading("Novelty / Innovation")
    doc.add_paragraph().add_run("What makes this project different?").bold = True
    add_justified_para("Unlike standard charting applications, this system synthesizes multiple domains of quantitative finance—Deep Learning forecasting, traditional Technical Indicators, and Natural Language Processing (NLP) sentiment—into a single integer score for actionable trading.")

    doc.add_paragraph().add_run("What value does it add?").bold = True
    doc.add_paragraph("Eliminates the need to switch between charting apps and news aggregators.", style='List Bullet')
    doc.add_paragraph("Reduces emotional bias in trading by relying on mathematical models.", style='List Bullet')
    doc.add_paragraph("Democratizes access to algorithmic trading techniques.", style='List Bullet')

    doc.add_paragraph().add_run("How does it improve existing solutions?").bold = True
    cmp_table = doc.add_table(rows=6, cols=2)
    cmp_table.style = 'Table Grid'
    hdr_cells_cmp = cmp_table.rows[0].cells
    hdr_cells_cmp[0].text = 'Existing Systems'
    hdr_cells_cmp[1].text = 'Proposed System (Stock Price Prediction)'
    for cell in hdr_cells_cmp:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    
    comparisons = [
        ("Manual Technical Analysis", "Automated vectorized calculation of indicators"),
        ("Basic LSTM forecasting", "Meta-Ensemble (LSTM + Attention + GRU + Conv1D)"),
        ("Isolated Price Data", "Integrated FinBERT NLP News Sentiment"),
        ("High-latency ML training", "Local File System Cache Strategy for instant inference"),
        ("Generic chart plotting", "Interactive Plotly graphs with dual-axis handling")
    ]
    for i, (old, new) in enumerate(comparisons, 1):
        cmp_table.rows[i].cells[0].text = old
        cmp_table.rows[i].cells[1].text = new
        for cell in cmp_table.rows[i].cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    add_section_heading("Technologies Used")
    doc.add_paragraph("Frontend: Streamlit, Plotly, HTML/CSS", style='List Bullet')
    doc.add_paragraph("Backend: Flask (Python), Flask-CORS", style='List Bullet')
    doc.add_paragraph("Machine Learning / AI: TensorFlow, Keras, Scikit-learn", style='List Bullet')
    doc.add_paragraph("Data Processing & Analytics: Pandas, NumPy", style='List Bullet')
    doc.add_paragraph("NLP: FinBERT (Transformers)", style='List Bullet')
    doc.add_paragraph("APIs: yfinance", style='List Bullet')

    doc.add_paragraph()

    add_section_heading("Expected Outcome")
    add_justified_para("The project aims to provide retail investors and traders with a comprehensive, adaptive, and highly accurate financial analysis platform. By fusing artificial intelligence with traditional market indicators, the platform will help users make smarter, data-driven investment decisions while reducing risk and enhancing their overall market understanding.")

    doc.add_paragraph()
    doc.add_paragraph().add_run("Team Members").bold = True
    doc.add_paragraph("[Name 1] ([ID 1]). _________________")
    doc.add_paragraph("[Name 2] ([ID 2]). _________________")
    doc.add_paragraph("[Name 3] ([ID 3]). _________________")
    doc.add_paragraph("[Name 4] ([ID 4]). _________________")

    doc.add_paragraph()
    doc.add_paragraph().add_run("Project Guide & Co-ordinator").bold = True
    doc.add_paragraph("[Name of Project Guide]. _________________")

    file_path = os.path.join(os.getcwd(), 'Project_Proposal_Report_v6.docx')
    doc.save(file_path)
    print(f"Successfully saved to {file_path}")

if __name__ == "__main__":
    create_document()
